"""Excel 导出 / 导入（台账导出、上线初始化导入、ERP 接口不通时的兜底导入）。"""
import io
from urllib.parse import quote

from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def export(title: str, columns: list[tuple[str, str]], rows: list[dict]) -> StreamingResponse:
    """columns: [(字段, 表头)]"""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append([c[1] for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDE7F3")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([_v(r.get(c[0])) for c in columns])
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(40, len(c[1]) * 2 + 4))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(title)}.xlsx"})


def _v(v):
    if isinstance(v, (list, tuple)):
        return "、".join(str(x) for x in v)
    if isinstance(v, bool):
        return "是" if v else "否"
    return v


def read_rows(content: bytes, mapping: dict[str, str]) -> list[dict]:
    """按表头读取：mapping {表头中文: 字段}"""
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    head = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {mapping[h]: i for i, h in enumerate(head) if h in mapping}
    out = []
    for r in rows[1:]:
        if not any(r):
            continue
        out.append({f: r[i] for f, i in idx.items()})
    return out
